"""
Markdown to DOCX Converter with README Title Support
Version: 2.0
Author: Codex
Date: 2026-08-05
Description:
    Reads a Markdown README file, extracts the backticked read-order titles,
    resolves each title to a file in the same folder, converts Markdown files
    to Word content, and adds a simple placeholder for non-Markdown sources.
Usage:
    python md_to_docx_readme_titles.py <root_folder> <readme_path> <output_docx>
"""

from __future__ import annotations

import difflib
import re
import sys
from pathlib import Path

import markdown
from bs4 import BeautifulSoup
from docx import Document


def markdown_to_html_soup(md_content: str) -> BeautifulSoup:
    """Convert Markdown text to a BeautifulSoup-parsable HTML structure."""
    html = markdown.markdown(md_content, extensions=["tables"])
    return BeautifulSoup(html, "html.parser")


def add_html_to_docx(soup: BeautifulSoup, doc: Document) -> None:
    """Add parsed HTML content to a Word document object."""
    for elem in soup.contents:
        if not getattr(elem, "name", None):
            continue
        if re.fullmatch(r"h[1-6]", elem.name):
            level = int(elem.name[1])
            doc.add_heading(elem.get_text(), level=level)
        elif elem.name == "hr":
            doc.add_paragraph()
        elif elem.name == "p":
            doc.add_paragraph(elem.get_text())
        elif elem.name in ["ul", "ol"]:
            style = "List Bullet" if elem.name == "ul" else "List Number"
            for li in elem.find_all("li", recursive=False):
                doc.add_paragraph(li.get_text(), style=style)
        elif elem.name == "blockquote":
            doc.add_paragraph(elem.get_text(), style="Intense Quote")
        elif elem.name == "table":
            rows = elem.find_all("tr")
            if rows:
                cols = rows[0].find_all(["td", "th"])
                table = doc.add_table(rows=len(rows), cols=len(cols))
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row.find_all(["td", "th"])):
                        table.cell(i, j).text = cell.get_text()


def normalize_loose(text: str) -> str:
    """Normalize text for fuzzy matching while keeping word boundaries."""
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", text.lower())).strip()


def normalize_compact(text: str) -> str:
    """Normalize text for strict matching by stripping punctuation and spaces."""
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def read_first_heading(path: Path) -> str:
    """Read the first Markdown heading from a file."""
    try:
        with path.open("r", encoding="utf-8") as f:
            for line in f:
                if line.lstrip().startswith("#"):
                    return line.lstrip("#").strip()
    except UnicodeDecodeError:
        pass
    return path.stem.replace("_", " ").strip()


def extract_read_order_entries(readme_path: Path) -> list[str]:
    """Extract backticked read-order entries from the README."""
    lines = readme_path.read_text(encoding="utf-8").splitlines()
    entries: list[str] = []
    in_read_order = False

    for line in lines:
        stripped = line.strip()
        if stripped.lower() == "## read order":
            in_read_order = True
            continue
        if in_read_order:
            if stripped.startswith("## "):
                break
            match = re.match(r'^\s*\d+\.\s+`([^`]+)`', line)
            if match:
                entries.append(match.group(1))
    return entries


def build_markdown_index(root_folder: Path) -> list[dict[str, str | Path]]:
    """Build a simple index of Markdown files in the root folder."""
    index: list[dict[str, str | Path]] = []
    for md_path in sorted(root_folder.glob("*.md")):
        if md_path.name.lower() == "readme.md":
            continue
        heading = read_first_heading(md_path)
        index.append(
            {
                "path": md_path,
                "heading": heading,
                "heading_loose": normalize_loose(heading),
                "heading_compact": normalize_compact(heading),
                "stem_loose": normalize_loose(md_path.stem),
                "stem_compact": normalize_compact(md_path.stem),
                "name_loose": normalize_loose(md_path.name),
                "name_compact": normalize_compact(md_path.name),
            }
        )
    return index


def score_entry(entry: str, info: dict[str, str | Path]) -> float:
    """Score how well a README entry matches a Markdown file."""
    entry_loose = normalize_loose(entry)
    entry_compact = normalize_compact(entry)

    if entry_loose in {info["heading_loose"], info["stem_loose"], info["name_loose"]}:
        return 1.0
    if entry_compact in {info["heading_compact"], info["stem_compact"], info["name_compact"]}:
        return 1.0

    scores = [
        difflib.SequenceMatcher(None, entry_loose, info["heading_loose"]).ratio(),
        difflib.SequenceMatcher(None, entry_loose, info["stem_loose"]).ratio(),
        difflib.SequenceMatcher(None, entry_loose, info["name_loose"]).ratio(),
        difflib.SequenceMatcher(None, entry_compact, info["heading_compact"]).ratio(),
        difflib.SequenceMatcher(None, entry_compact, info["stem_compact"]).ratio(),
        difflib.SequenceMatcher(None, entry_compact, info["name_compact"]).ratio(),
    ]
    return max(scores)


def resolve_entry_to_path(entry: str, root_folder: Path, md_index: list[dict[str, str | Path]]) -> Path | None:
    """Resolve a README entry to a file path."""
    candidate = root_folder / entry
    if candidate.exists():
        return candidate

    candidate = root_folder / Path(entry).name
    if candidate.exists():
        return candidate

    best_path: Path | None = None
    best_score = -1.0
    for info in md_index:
        score = score_entry(entry, info)
        if score > best_score:
            best_score = score
            best_path = info["path"]  # type: ignore[assignment]

    if best_score >= 0.60:
        return best_path
    return None


def add_source_placeholder(doc: Document, source_name: str, note: str) -> None:
    """Add a placeholder section for a source that is not Markdown."""
    doc.add_heading(source_name, level=1)
    doc.add_paragraph(note)


def create_word_document(root_folder: str, readme_path: str, output_path: str) -> None:
    """Create a Word document from the README read order."""
    root = Path(root_folder)
    readme = Path(readme_path)
    doc = Document()

    readme_title = read_first_heading(readme)
    doc.add_heading(readme_title or "Consolidated Markdown Documents", 0)

    md_index = build_markdown_index(root)
    entries = extract_read_order_entries(readme)

    if not entries:
        raise ValueError(f"No backticked read-order entries found in {readme}")

    for i, entry in enumerate(entries):
        source_path = resolve_entry_to_path(entry, root, md_index)
        if i > 0:
            doc.add_page_break()

        if source_path is None:
            add_source_placeholder(doc, entry, "No matching file was found in the folder.")
            print(f"WARNING: No match found for: {entry}")
            continue

        suffix = source_path.suffix.lower()
        if suffix == ".md":
            with source_path.open("r", encoding="utf-8") as f:
                md = f.read()
            soup = markdown_to_html_soup(md)
            add_html_to_docx(soup, doc)
        else:
            add_source_placeholder(
                doc,
                source_path.name,
                "This source file is not converted by this script.",
            )
            print(f"INFO: Non-Markdown source noted: {source_path}")

    doc.save(output_path)
    print(f"✅ Document saved to: {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("Usage: python md_to_docx_readme_titles.py <root_folder> <readme_path> <output_docx>")
        sys.exit(1)

    root_folder = sys.argv[1]
    readme_path = sys.argv[2]
    output_path = sys.argv[3]

    create_word_document(root_folder, readme_path, output_path)
