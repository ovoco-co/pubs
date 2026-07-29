"""
Markdown to DOCX Converter with README Link Support
Version: 1.1
Author: Craig Wilson
Date: 2025-02-20
Description:
    This script reads a Markdown README file, extracts links to other markdown documents,
    converts each to HTML using the `markdown` module with table support,
    then parses the HTML using BeautifulSoup to construct a Word document
    using python-docx. The output is a consolidated DOCX representing all linked content.

Dependencies:
    - markdown
    - beautifulsoup4
    - python-docx
Usage:
    python md_to_docx.py <root_folder> <readme_path> <output_docx>

Example:
    python md_to_docx.py ./docs ./docs/README.md ./output/architecture.docx
"""

import os
import sys
import re
import markdown
from bs4 import BeautifulSoup
from docx import Document

def markdown_to_html_soup(md_content):
    """
    Convert Markdown text to a BeautifulSoup-parsable HTML structure.

    Args:
        md_content (str): The raw Markdown content.

    Returns:
        BeautifulSoup: Parsed HTML content as a soup object.
    """
    html = markdown.markdown(md_content, extensions=['tables'])
    return BeautifulSoup(html, "html.parser")

def add_html_to_docx(soup, doc):
    """
    Add parsed HTML content to a Word document object.

    Args:
        soup (BeautifulSoup): Parsed HTML content.
        doc (Document): The python-docx Document object to append to.
    """
    for elem in soup.contents:
        if elem.name and re.fullmatch(r'h[1-6]', elem.name):
            level = int(elem.name[1])
            doc.add_heading(elem.get_text(), level=level)
        elif elem.name == 'hr':
            doc.add_paragraph()
        elif elem.name == 'p':
            doc.add_paragraph(elem.get_text())
        elif elem.name in ['ul', 'ol']:
            style = 'List Bullet' if elem.name == 'ul' else 'List Number'
            for li in elem.find_all('li'):
                doc.add_paragraph(li.get_text(), style=style)
        elif elem.name == 'blockquote':
            doc.add_paragraph(elem.get_text(), style='Intense Quote')
        elif elem.name == 'table':
            rows = elem.find_all('tr')
            if rows:
                cols = rows[0].find_all(['td', 'th'])
                table = doc.add_table(rows=len(rows), cols=len(cols))
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row.find_all(['td', 'th'])):
                        table.cell(i, j).text = cell.get_text()

def extract_links_from_readme(readme_path):
    """
    Extract Markdown-style links from a README file.

    Args:
        readme_path (str): Path to the README file.

    Returns:
        list[str]: List of linked file paths.
    """
    with open(readme_path, 'r', encoding='utf-8') as f:
        content = f.read()
    return re.findall(r'\[.*?\]\((.*?)\)', content)

def create_word_document(root_folder, readme_path, output_path):
    """
    Create a Word document by reading Markdown files referenced in the README.

    Args:
        root_folder (str): Base folder to resolve relative Markdown paths.
        readme_path (str): Path to the README markdown file.
        output_path (str): Destination path for the output DOCX file.
    """
    doc = Document()
    doc.add_heading("Azure Landing Zone Cybersecurity Architecture", 0)

    links = extract_links_from_readme(readme_path)

    for link in links:
        full_path = os.path.join(root_folder, link)
        if os.path.exists(full_path):
            with open(full_path, 'r', encoding='utf-8') as f:
                md = f.read()
            soup = markdown_to_html_soup(md)
            doc.add_page_break()
            add_html_to_docx(soup, doc)
        else:
            print(f"WARNING: File not found - {full_path}")

    doc.save(output_path)
    print(f"✅ Document saved to: {output_path}")

if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("Usage: python md_to_docx.py <root_folder> <readme_path> <output_docx>")
        sys.exit(1)

    root_folder = sys.argv[1]
    readme_path = sys.argv[2]
    output_path = sys.argv[3]

    create_word_document(root_folder, readme_path, output_path)
